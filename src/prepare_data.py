"""Stage 1 only: preserve, normalize and reconcile source sales-detail lines.

The workbook is authoritative. The publisher CSV is used only after preparation
for a positional comparison; it supplies no identifiers or analytical values.
"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
import hashlib
import json
from pathlib import Path
import platform
import re
import sys
import xml.etree.ElementTree as ET
import zipfile

import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
SHEET = "Sales Details 2010 - 18122015"
REFERENCE_SHA256 = "80e329dcc1ac78bf572440f876a4a3884145165ce79be2d3c6c8510312fcdd37"
CENT = Decimal("0.01")
HEADERS = ["#", "Coil Pattern", "Group Name1", "Product", "Document Number",
           "Doc. Date", "Week", "Month", "Group Name2", "Customer Code",
           "Customer Name", "Item Code", "Item Description", "Thickness",
           "Width", "Colour", "Batch Number", "ItemCost", "Sales Price",
           "Batch Quantity", "Vat Value", "Sales Amount"]
RENAME = dict(zip(HEADERS, ["source_line_number", "coil_pattern", "product_group",
    "product", "document_number", "source_document_date", "source_week_label",
    "source_month", "customer_group", "customer_code", "customer_name", "item_code",
    "item_description", "thickness_code", "width", "source_colour", "batch_number",
    "item_cost", "sales_price", "reported_quantity", "vat_value", "reported_sales_amount"]))
IDENTIFIERS = {"source_line_number", "document_number", "customer_code", "item_code", "batch_number"}
NUMERICS = {"width", "item_cost", "sales_price", "reported_quantity", "vat_value", "reported_sales_amount"}
COLOUR_ALIASES = {"plain": "Plain", "H.B GREEN": "H.B Green"}


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def decimal_value(value):
    if value is None or isinstance(value, bool):
        raise ValueError(f"Missing or invalid numeric value: {value!r}")
    number = Decimal(str(value))
    if not number.is_finite():
        raise ValueError(f"Non-finite numeric value: {value!r}")
    return number


def identifier_text(value, number_format="General"):
    """Keep text identifiers, including zeroes. Numeric zero-padding only if explicit."""
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    number = decimal_value(value)
    if number != number.to_integral_value():
        raise ValueError(f"Nonintegral identifier: {value!r}")
    label = str(int(number))
    if re.fullmatch(r"0+", number_format):
        label = label.zfill(len(number_format))
    elif number_format != "General":
        raise ValueError(f"Unsupported numeric identifier format: {number_format!r}")
    return label


def text_value(value):
    return "" if value is None else str(value).strip()


def normalize_date(value, source_month):
    """Return (date or None, corrected, rule, issue); never guess unresolved dates."""
    try:
        month_decimal = decimal_value(source_month)
        if month_decimal != month_decimal.to_integral_value() or not 1 <= month_decimal <= 12:
            raise ValueError("Source month must be an integer from 1 to 12")
        month = int(month_decimal)
        if isinstance(value, datetime):
            if value.time() != datetime.min.time():
                raise ValueError("Unexpected time component")
            original = value.date()
            if original.month == month:
                return original, False, "datetime_month_matches", ""
            if original.day == month:
                corrected = date(original.year, original.day, original.month)
                return corrected, True, "datetime_day_month_swapped", ""
            raise ValueError("Neither decoded month nor day matches source month")
        if isinstance(value, str) and re.fullmatch(r"\d{2}-\d{2}-\d{4}", value.strip()):
            parsed = datetime.strptime(value.strip(), "%d-%m-%Y").date()
            if parsed.month != month:
                raise ValueError("Day-first string month disagrees with source month")
            return parsed, False, "string_day_first", ""
        raise ValueError("Unexpected date representation")
    except (ValueError, InvalidOperation, TypeError) as exc:
        return None, False, "unresolved", str(exc)


def xml_source_controls(path):
    """Independent controls from stored worksheet XML, without pandas/openpyxl values."""
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    amount = Decimal(0)
    rows = 0
    values = 0
    with zipfile.ZipFile(path) as archive, archive.open("xl/worksheets/sheet1.xml") as stream:
        for _, element in ET.iterparse(stream, events=["end"]):
            if element.tag == f"{{{ns['s']}}}row":
                if int(element.attrib["r"]) > 1:
                    rows += 1
                    for cell in element.findall("s:c", ns):
                        if re.fullmatch(r"V\d+", cell.attrib["r"]):
                            v = cell.find("s:v", ns)
                            if v is not None:
                                amount += Decimal(v.text)
                                values += 1
                element.clear()
    return {"rows": rows, "sales_values": values, "sales_raw_xml": str(amount),
            "sales_at_cent_precision": str(amount.quantize(CENT))}


def write_csv(frame, path):
    frame.to_csv(path, index=False, lineterminator="\n")


def run(args):
    processed = args.output_root / "data/processed"
    reports = args.output_root / "reports"
    processed.mkdir(parents=True, exist_ok=True)
    reports.mkdir(parents=True, exist_ok=True)
    workbook_hash = sha256(args.workbook)
    if workbook_hash != REFERENCE_SHA256:
        raise ValueError("Workbook checksum differs from the reviewed source. Reassess before proceeding.")
    workbook = openpyxl.load_workbook(args.workbook, read_only=True, data_only=False)
    if workbook.sheetnames != [SHEET]:
        raise ValueError("Unexpected workbook worksheets")
    sheet = workbook[SHEET]
    rows = sheet.iter_rows()
    header = [c.value for c in next(rows)]
    if header[:22] != HEADERS or header[22:] != [None, None, None]:
        raise ValueError("Unexpected source schema")
    records, raw_values, issues = [], [], []
    source_types, source_formats = defaultdict(Counter), defaultdict(Counter)
    missing, trims = Counter(), Counter()
    formulas = trailing_values = 0
    source_total = Decimal(0)
    for row_id, cells in enumerate(rows, start=2):
        trailing_values += sum(c.value is not None for c in cells[22:])
        values = [c.value for c in cells[:22]]
        raw_values.append(values)
        record = {"source_row_id": row_id}
        for source, cell in zip(HEADERS, cells[:22]):
            value = cell.value
            source_types[source][type(value).__name__] += 1
            source_formats[source][getattr(cell, "number_format", "")] += 1
            missing[source] += value is None
            trims[source] += isinstance(value, str) and value != value.strip()
            formulas += cell.data_type == "f"
            target = RENAME[source]
            if target == "source_document_date":
                record[target] = value.isoformat() if isinstance(value, datetime) else text_value(value)
                record["source_date_kind"] = "excel_datetime" if isinstance(value, datetime) else "text"
                record["source_date_number_format"] = cell.number_format
            elif target in IDENTIFIERS:
                record[target] = identifier_text(value, cell.number_format)
            elif target in NUMERICS:
                record[target] = str(decimal_value(value))
            elif target == "source_month":
                record[target] = text_value(value)
            else:
                record[target] = text_value(value)
        normalized, corrected, rule, issue = normalize_date(values[5], values[7])
        record.update(document_date=normalized.isoformat() if normalized else "",
                      document_year=normalized.year if normalized else "",
                      document_month=normalized.month if normalized else "",
                      document_day=normalized.day if normalized else "",
                      is_primary_window=bool(normalized and 2011 <= normalized.year <= 2014),
                      date_corrected=corrected, date_normalization_rule=rule, date_issue=issue,
                      colour=COLOUR_ALIASES.get(record["source_colour"], record["source_colour"]))
        if issue:
            issues.append({"source_row_id": row_id, "issue": issue})
        source_total += decimal_value(values[21])
        records.append(record)
    workbook.close()
    frame = pd.DataFrame(records)
    raw = pd.DataFrame(raw_values, columns=HEADERS)
    dated = frame.sort_values(["document_date", "source_row_id"])
    display_names = dated.drop_duplicates("customer_code", keep="last").set_index("customer_code")["customer_name"]
    frame["customer_display_name"] = frame["customer_code"].map(display_names)
    output = processed / "sales_lines.csv"
    write_csv(frame, output)
    # Re-read serialized bytes; identifiers intentionally remain strings.
    reread = pd.read_csv(output, dtype=str, keep_default_na=False)
    clean_total = sum((Decimal(x) for x in reread.reported_sales_amount), Decimal(0))
    xml = xml_source_controls(args.workbook)
    checks = []
    def check(name, actual, expected, passed, severity="critical"):
        checks.append(dict(check_name=name, actual_value=str(actual), expected_rule=str(expected),
                           status="PASS" if passed else "FAIL", severity=severity))
    check("source_hash_matches_reviewed_workbook", workbook_hash, REFERENCE_SHA256, workbook_hash == REFERENCE_SHA256)
    check("source_to_prepared_rows", len(reread), len(raw), len(reread) == len(raw))
    check("xml_to_reader_rows", len(raw), xml["rows"], len(raw) == xml["rows"])
    check("xml_sales_cell_count", xml["sales_values"], len(raw), xml["sales_values"] == len(raw))
    check("source_to_prepared_sales", clean_total, source_total, clean_total == source_total)
    xml_delta = abs(Decimal(xml["sales_raw_xml"]) - source_total)
    check("xml_to_reader_sales", xml_delta, "absolute difference <= 0.01 source currency units", xml_delta <= CENT)
    check("source_row_ids_unique", reread.source_row_id.nunique(), len(raw), reread.source_row_id.is_unique)
    check("source_row_ids_preserved", int(reread.source_row_id.eq(frame.source_row_id.astype(str)).sum()), len(raw), reread.source_row_id.tolist() == frame.source_row_id.astype(str).tolist())
    check("empty_trailing_columns", trailing_values, 0, trailing_values == 0)
    check("no_source_formulas", formulas, 0, formulas == 0)
    check("required_dates_resolved", len(issues), 0, len(issues) == 0)
    month_bad = sum(bool(r.document_date) and int(r.document_month) != int(r.source_month) for r in frame.itertuples())
    check("normalized_date_matches_source_month", month_bad, 0, month_bad == 0)
    for c in ["customer_code", "item_code", "batch_number", "document_number"]:
        check(f"{c}_not_blank", int(reread[c].eq("").sum()), 0, not reread[c].eq("").any())
        expected = [identifier_text(v, "General") for v in raw[HEADERS[list(RENAME.values()).index(c)]]]
        check(f"{c}_preserved", int(reread[c].eq(pd.Series(expected)).sum()), len(raw), reread[c].tolist() == expected)
    residuals = [abs(Decimal(s) - Decimal(p) * Decimal(q)) for s,p,q in zip(frame.reported_sales_amount, frame.sales_price, frame.reported_quantity)]
    check("sales_price_quantity_arithmetic", sum(v > CENT for v in residuals), "0 lines with absolute difference > 0.01", all(v <= CENT for v in residuals))
    suffix_raw = raw.Thickness.map(lambda v: isinstance(v, str) and bool(re.search(r"[A-Za-z]", v)))
    suffix_kept = sum(reread.loc[suffix_raw, "thickness_code"].eq(raw.loc[suffix_raw, "Thickness"]).tolist())
    check("thickness_suffixes_preserved", suffix_kept, int(suffix_raw.sum()), suffix_kept == int(suffix_raw.sum()))
    primary = frame.loc[frame.is_primary_window]
    primary_months = primary[["document_year", "document_month"]].drop_duplicates()
    check("primary_calendar_months_observed", len(primary_months), "48 (2011–2014); not a completeness guarantee", len(primary_months) == 48)
    for c in sorted(NUMERICS):
        check(f"{c}_nonnegative", sum(Decimal(v) < 0 for v in frame[c]), 0, all(Decimal(v) >= 0 for v in frame[c]))
    check("reported_quantity_positive", sum(Decimal(v) <= 0 for v in frame.reported_quantity), 0, all(Decimal(v) > 0 for v in frame.reported_quantity))

    month_rows = []
    for (year, month), group in frame.groupby(["document_year", "document_month"]):
        month_rows.append(dict(year=year, month=month, source_lines=len(group), customers=group.customer_code.nunique(),
            items=group.item_code.nunique(), first_recorded_date=group.document_date.min(),
            last_recorded_date=group.document_date.max(), reported_sales_amount=str(sum((Decimal(x) for x in group.reported_sales_amount), Decimal(0))),
            date_corrections=int(group.date_corrected.sum()), is_primary_window=bool(group.is_primary_window.iloc[0])))
    write_csv(pd.DataFrame(month_rows), reports / "source_month_coverage.csv")
    def variant_report(key, label, filename):
        counts = frame.groupby(key)[label].nunique()
        subset = frame.loc[frame[key].isin(counts.loc[counts > 1].index)]
        summary = subset.groupby([key,label]).agg(source_lines=("source_row_id", "size"),
            first_recorded_date=("document_date", "min"), last_recorded_date=("document_date", "max"),
            reported_sales_amount=("reported_sales_amount", lambda x:str(sum((Decimal(v) for v in x), Decimal(0))))).reset_index()
        write_csv(summary, reports / filename)
        return int((counts > 1).sum())
    name_variant_count = variant_report("customer_code", "customer_name", "customer_name_variants.csv")
    description_variant_count = variant_report("item_code", "item_description", "item_description_variants.csv")
    zero_sales = frame.loc[frame.reported_sales_amount.map(Decimal).eq(0)].copy()
    zero_sales["audit_treatment"] = zero_sales.sales_price.map(lambda v: "retained_zero_price" if Decimal(v) == 0 else "retained_product_below_half_cent")
    write_csv(zero_sales, reports / "zero_sales_lines.csv")

    comparison_checks, differences = [], []
    comparison_sha = None
    if args.comparison.exists():
        comparison_sha = sha256(args.comparison)
        comparison = pd.read_csv(args.comparison, dtype=str, keep_default_na=False)
        csv_headers = ["coil_pattern","grp_name","thickness","width","colour","item_Cost","sales_price","vatval","year","month","day","week","order_qty"]
        aligned = len(comparison) == len(frame) and comparison.columns.tolist() == csv_headers
        comparison_checks.append(dict(field="rows_and_schema", compared_lines=len(comparison), differences=0 if aligned else 1,
            status="PASS" if aligned else "WARNING", comparison_rule="Same row count and expected 13-column schema"))
        if aligned:
            mapping = {"coil_pattern":"coil_pattern", "grp_name":"customer_group", "thickness":"thickness_code",
                "width":"width", "colour":"source_colour", "item_Cost":"item_cost", "sales_price":"sales_price",
                "vatval":"vat_value", "year":"document_year", "month":"document_month", "day":"document_day", "order_qty":"reported_quantity"}
            numeric_fields = {"thickness","width","item_Cost","sales_price","vatval","year","month","day","order_qty"}
            for field in csv_headers:
                bad = 0
                for i, value in enumerate(comparison[field]):
                    if field == "week":
                        original = str(date.fromisoformat(frame.at[i,"document_date"]).weekday()) if frame.at[i,"document_date"] else ""
                        matches = original == value
                    else:
                        original = str(frame.at[i, mapping[field]])
                        if field in numeric_fields:
                            try:
                                matches = Decimal(original) == Decimal(value)
                            except InvalidOperation:
                                matches = original == value
                        else:
                            matches = original == value
                    if not matches:
                        bad += 1
                        differences.append(dict(source_row_id=frame.at[i,"source_row_id"], comparison_field=field,
                            prepared_source_value=original, publisher_csv_value=value))
                comparison_checks.append(dict(field=field, compared_lines=len(frame), differences=bad,
                    status="PASS" if bad == 0 else "WARNING", comparison_rule="CSV week compared with date weekday (Monday=0)" if field == "week" else "Positional equality; numerics compared as Decimal"))
    else:
        comparison_checks.append(dict(field="comparison_file", compared_lines=0, differences="", status="NOT_RUN", comparison_rule="Optional comparison CSV absent"))
    write_csv(pd.DataFrame(comparison_checks), reports / "comparison_checks.csv")
    write_csv(pd.DataFrame(differences, columns=["source_row_id","comparison_field","prepared_source_value","publisher_csv_value"]), reports / "comparison_differences.csv")
    write_csv(pd.DataFrame(checks), reports / "preparation_checks.csv")

    batch_cases = frame.groupby(frame.batch_number.str.casefold()).batch_number.agg(lambda v:sorted(set(v)))
    batch_cases = batch_cases.loc[batch_cases.map(len) > 1].to_dict()
    item_variation = {c:int(frame.groupby("item_code")[c].nunique().gt(1).sum()) for c in ["thickness_code","width","source_colour","coil_pattern","item_description"]}
    numeric_profile = {c:dict(minimum=str(min(map(Decimal, frame[c]))), maximum=str(max(map(Decimal,frame[c]))),
        zero_lines=sum(Decimal(v)==0 for v in frame[c]), negative_lines=sum(Decimal(v)<0 for v in frame[c])) for c in sorted(NUMERICS)}
    audit = dict(stage=1, pipeline_status="PASS" if all(c["status"] == "PASS" for c in checks) else "FAIL",
        source=dict(filename=args.workbook.name, sha256=workbook_hash, worksheet=SHEET,
            worksheet_rows_including_header=sheet.max_row, data_rows=len(raw), columns=sheet.max_column,
            populated_columns=22, trailing_nonempty_cells=trailing_values, formula_cells=formulas,
            types={k:dict(v) for k,v in source_types.items()}, number_formats={k:dict(v) for k,v in source_formats.items()},
            missing_values=dict(missing), whitespace_trim_counts=dict(trims)),
        runtime=dict(python=platform.python_version(), pandas=pd.__version__, openpyxl=openpyxl.__version__),
        counts=dict(prepared_rows=len(reread), excluded_rows=0, customer_codes=frame.customer_code.nunique(),
            item_codes=frame.item_code.nunique(), document_numbers=frame.document_number.nunique(), batch_numbers=frame.batch_number.nunique(),
            primary_rows=len(primary), primary_customer_codes=primary.customer_code.nunique(), primary_item_codes=primary.item_code.nunique(),
            primary_calendar_months=len(primary_months), customer_name_variant_codes=name_variant_count,
            item_description_variant_codes=description_variant_count, zero_sales_lines=len(zero_sales)),
        reconciliation=dict(source_sales=str(source_total), prepared_sales=str(clean_total), source_prepared_difference=str(clean_total-source_total),
            xml_control=xml, xml_reader_absolute_difference=str(xml_delta), sales_price_quantity_max_absolute_difference=str(max(residuals)),
            sales_price_quantity_nonexact_lines=sum(v!=0 for v in residuals), sales_price_quantity_over_cent_lines=sum(v>CENT for v in residuals)),
        dates=dict(first=frame.document_date.min(), last=frame.document_date.max(), normalization_counts=frame.date_normalization_rule.value_counts().to_dict(),
            correction_count=int(frame.date_corrected.sum()), unresolved=issues),
        duplicates=dict(exact_source_rows=int(raw.duplicated(keep=False).sum()),
            all_fields_except_source_sequence=int(raw.drop(columns="#").duplicated(keep=False).sum()),
            all_fields_except_sequence_and_batch=int(raw.drop(columns=["#","Batch Number"]).duplicated(keep=False).sum()),
            interpretation="No rows removed. Weaker-key similarities can be distinct batches."),
        documents=dict(numbers_with_multiple_dates=int(frame.groupby("document_number").document_date.nunique().gt(1).sum()),
            year_number_with_multiple_dates=int(frame.groupby(["document_year","document_number"]).document_date.nunique().gt(1).sum()),
            interpretation="Numbers repeat across years; neither number alone nor year+number is certified as an order key."),
        dimensions={c:frame[c].value_counts().to_dict() for c in ["product","product_group","thickness_code","width","source_colour","colour","coil_pattern","customer_group"]},
        item_specification_variation=item_variation, numeric_profile=numeric_profile,
        colour_alias_changes={k:int(frame.source_colour.eq(k).sum()) for k in COLOUR_ALIASES}, batch_case_variants=batch_cases,
        comparison=dict(filename=args.comparison.name, sha256=comparison_sha, basis="Row-position comparison, supported by aligned shared fields; not independent date proof.", checks=comparison_checks),
        outputs=dict(prepared_csv_sha256=sha256(output), preparation_checks=len(checks)),
        unresolved_issues=[
            dict(issue="Currency, quantity units, cost basis and net/gross sales semantics are not documented", treatment="Use reported sales in source currency units; no margin or combined cross-item quantity/price KPI. Item-level bridge requires an explicit consistent-unit assumption unless additional metadata confirms it."),
            dict(issue="One account code has sequential name labels; legal rename not verified", treatment="Group by customer_code. Keep original names; latest observed name is display-only."),
            dict(issue="Fourteen item codes have multiple descriptions, including conflicting specification wording", treatment="Keep original descriptions and stable structured specifications. Item-code identity does not independently verify quantity-unit stability. Carry the issue into bridge eligibility review."),
            dict(issue="Zero-sales purpose and very low prices are unverified", treatment="Retain all lines and their reported amounts and quantities. Do not relabel as gifts, returns, cancellations or data errors."),
            dict(issue="Document and batch keys lack business certification", treatment="Preserve exact identifiers and case. No order metrics, batch deduplication or unsupported joins."),
            dict(issue="True dates and transaction-population completeness lack independent evidence", treatment="Apply documented source-month normalization. Use 2011–2014 for annual comparisons, with observed-month coverage only."),
        ])
    (processed / "ingest_audit.json").write_text(json.dumps(audit, indent=2, sort_keys=True) + "\n")
    print(json.dumps({"stage":1, "status":audit["pipeline_status"], "rows":len(frame), "source_sales":str(source_total),
        "prepared_sales":str(clean_total), "date_corrections":int(frame.date_corrected.sum()), "unresolved_dates":len(issues),
        "primary_rows":len(primary), "checks_passed":sum(c["status"]=="PASS" for c in checks), "checks_total":len(checks)}, indent=2))
    if audit["pipeline_status"] != "PASS":
        return 1
    return 0


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=ROOT / "data/raw/aluminum_sales.xlsx")
    parser.add_argument("--comparison", type=Path, default=ROOT / "data/comparison/publisher_sales.csv")
    parser.add_argument("--output-root", type=Path, default=ROOT)
    sys.exit(run(parser.parse_args()))
